"""资料包导出 API（T09/T10/T11）。"""

import io
import zipfile
from pathlib import Path

import openpyxl
import pytest
from pypdf import PdfReader

from tests.conftest import FIXTURES_DIR
from tests.integration.test_batches_helpers import (
    add_items,
    batch_with,
    make_batch,
    make_expense,
    png_bytes,
    project_id,
    upload,
)

EXPORT_KEYS = {
    "id",
    "layout",
    "file_name",
    "url",
    "sha256",
    "item_count",
    "total_cents",
    "created_at",
    "created_by",
}


def export(client, batch_id: int, layout: str = "by_expense") -> dict:
    response = client.post(f"/api/batches/{batch_id}/export", json={"layout": layout})
    assert response.status_code == 200, response.text
    return response.json()["data"]


def download(client, record: dict) -> zipfile.ZipFile:
    response = client.get(record["url"])
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"
    return zipfile.ZipFile(io.BytesIO(response.content))


def summary_rows(archive: zipfile.ZipFile) -> tuple[list[tuple], list[tuple]]:
    workbook = openpyxl.load_workbook(io.BytesIO(archive.read("00_报销汇总表.xlsx")))
    assert workbook.sheetnames == ["汇总", "分类小计"]
    summary = list(workbook["汇总"].iter_rows(values_only=True))
    subtotal = list(workbook["分类小计"].iter_rows(values_only=True))
    assert workbook["汇总"].freeze_panes == "A3"
    return summary, subtotal


def cents(value) -> int:
    return round(value * 100)


@pytest.fixture
def eight_items(client):
    project = project_id(client, "科研A")
    batch = make_batch(client, "第1批", project_id=project)
    days = ["2026-09-03", "2026-09-01", "2026-09-02", "2026-09-01"] * 2
    expenses = [
        make_expense(
            client, spent_on=day, cents=12345 + i * 1001, merchant=f"商家{i}", project_id=project
        )
        for i, day in enumerate(days)
    ]
    upload(client, expenses[0]["id"], [("订单.png", png_bytes()), ("支付.png", png_bytes())])
    detail = add_items(client, batch["id"], [e["id"] for e in expenses]).json()["data"]
    return detail, expenses


def test_export_zip_structure_and_summary(client, settings, eight_items):  # T09
    batch, expenses = eight_items
    record = export(client, batch["id"])
    assert set(record) == EXPORT_KEYS
    assert record["item_count"] == 8
    assert record["total_cents"] == sum(e["amount_cents"] for e in expenses)
    assert record["file_name"].startswith("报销资料_科研A_") and record["url"].endswith("/file")

    archive = download(client, record)
    names = archive.namelist()
    assert names[:2] == ["00_报销汇总表.xlsx", "01_打印版_全部材料.pdf"]
    invoices = sorted(n for n in names if n.startswith("02_发票原件/"))
    assert len(invoices) == 8
    assert all(info.flag_bits & 0x800 for info in archive.infolist())  # UTF-8 文件名

    summary, subtotal = summary_rows(archive)
    assert "第1批" in summary[0][0] and "科研A" in summary[0][0]
    assert summary[1][:6] == ("序号", "日期", "商家", "摘要", "分类", "金额")
    body, total_row = summary[2:-1], summary[-1]
    assert [row[0] for row in body] == list(range(1, 9))
    assert [row[1] for row in body] == sorted(row[1] for row in body)
    assert sum(cents(row[5]) for row in body) == cents(total_row[5]) == record["total_cents"]
    assert total_row[0] == "合计"
    assert subtotal[-1] == ("合计", 8, total_row[5])
    assert subtotal[2][0] == "未分类"

    by_seq = {row[0]: row for row in body}
    for name in invoices:
        seq, category, amount, merchant = Path(name).stem.split("_")[:4]
        row = by_seq[int(seq)]
        assert (category, amount, merchant) == (row[4] or "未分类", f"{row[5]:.2f}", row[2])
    first = next(row for row in body if row[2] == "商家0")
    assert first[7] == "发票、订单明细、支付记录"
    group = f"03_支撑材料/{first[0]:02d}_商家0_{first[5]:.2f}/"
    assert {f"{group}订单明细_1.png", f"{group}支付记录_1.png"} <= set(names)

    pdf = PdfReader(io.BytesIO(archive.read("01_打印版_全部材料.pdf")))
    assert len(pdf.pages) == 10
    zip_path = settings.data_dir / "资料包"
    assert len(list(zip_path.glob("*_科研A_第1批/*.zip"))) == 1


def test_export_by_kind_layout(client, eight_items):
    batch, _ = eight_items
    record = export(client, batch["id"], "by_kind")
    assert record["layout"] == "by_kind"
    names = download(client, record).namelist()
    assert any(n.startswith("03_支撑材料/订单明细/") and n.endswith("_1.png") for n in names)
    assert any(n.startswith("03_支撑材料/支付记录/") for n in names)


def test_export_twice_keeps_old_package(client, settings, eight_items):  # T11
    batch, _ = eight_items
    first, second = export(client, batch["id"]), export(client, batch["id"])
    assert first["id"] != second["id"]
    assert first["file_name"] != second["file_name"]
    packages = list(settings.packages_dir.rglob("*.zip"))
    assert len(packages) == 2
    assert not list(settings.packages_dir.rglob("*.part"))
    detail = client.get(f"/api/batches/{batch['id']}").json()["data"]
    assert [item["id"] for item in detail["exports"]] == [second["id"], first["id"]]
    assert download(client, first).namelist()


def test_export_errors(client, settings):
    empty = make_batch(client)
    response = client.post(f"/api/batches/{empty['id']}/export", json={"layout": "by_expense"})
    assert response.status_code == 400
    assert response.json()["error"] == "空批次不能导出资料包"
    bad_layout = client.post(f"/api/batches/{empty['id']}/export", json={"layout": "x"})
    assert bad_layout.status_code == 422
    assert client.get("/api/exports/999/file").status_code == 404

    batch, _ = batch_with(client, 1)
    record = export(client, batch["id"])
    for path in settings.packages_dir.rglob("*.zip"):
        path.unlink()
    assert client.get(record["url"]).json()["error"] == "资料包文件不存在"


def test_export_skips_unmergeable_files_with_note(client):  # T10 附带：XML/损坏 PDF 不崩溃
    expense = make_expense(client, merchant="数电", note="原备注")
    xml = (FIXTURES_DIR / "invoices" / "digital_invoice.xml").read_bytes()
    corrupt = (FIXTURES_DIR / "invoices" / "corrupt.pdf").read_bytes()
    good = (FIXTURES_DIR / "invoices" / "vat_electronic_normal.pdf").read_bytes()
    upload(client, expense["id"], [("发票.xml", xml)], kind="invoice")
    upload(client, expense["id"], [("坏.pdf", corrupt)], kind="other")
    upload(client, expense["id"], [("好.pdf", good)], kind="order")
    batch = make_batch(client)
    add_items(client, batch["id"], [expense["id"]], force=True)

    archive = download(client, export(client, batch["id"]))
    summary, _ = summary_rows(archive)
    note = summary[2][9]
    assert note.startswith("原备注；")
    assert "发票.xml 未合并到打印版" in note and "坏.pdf 未合并到打印版" in note
    assert "好.pdf" not in note
    pdf = PdfReader(io.BytesIO(archive.read("01_打印版_全部材料.pdf")))
    good_pages = len(PdfReader(io.BytesIO(good)).pages)
    assert len(pdf.pages) == 1 + good_pages
    assert any(n.endswith(".xml") and n.startswith("02_发票原件/01_") for n in archive.namelist())


def test_export_skips_missing_source_file(client, settings):
    batch, expenses = batch_with(client, 1)
    attachment = expense_detail_first_attachment(client, expenses[0]["id"])
    library_file = next(settings.library_dir.rglob(attachment["file_name"]))
    library_file.unlink()
    archive = download(client, export(client, batch["id"]))
    assert not [n for n in archive.namelist() if n.startswith("02_发票原件/")]
    summary, _ = summary_rows(archive)
    assert "发票.png 未合并到打印版" in summary[2][9]


def expense_detail_first_attachment(client, expense_id: int) -> dict:
    return client.get(f"/api/expenses/{expense_id}").json()["data"]["attachments"][0]


def test_delete_export_removes_file_and_record(client, settings, eight_items):
    batch, _ = eight_items
    first, second = export(client, batch["id"]), export(client, batch["id"])

    response = client.delete(f"/api/exports/{first['id']}")

    assert response.status_code == 200 and response.json()["data"] is None
    assert len(list(settings.packages_dir.rglob("*.zip"))) == 1
    detail = client.get(f"/api/batches/{batch['id']}").json()["data"]
    assert [item["id"] for item in detail["exports"]] == [second["id"]]
    assert client.get(f"/api/exports/{first['id']}/file").status_code == 404


def test_delete_last_export_prunes_empty_package_dir(client, settings, eight_items):
    batch, _ = eight_items
    record = export(client, batch["id"])

    client.delete(f"/api/exports/{record['id']}")

    assert not [path for path in settings.packages_dir.iterdir() if path.is_dir()]


def test_delete_export_tolerates_missing_file_and_unknown_id(client, settings, eight_items):
    batch, _ = eight_items
    record = export(client, batch["id"])
    for path in settings.packages_dir.rglob("*.zip"):
        path.unlink()

    assert client.delete(f"/api/exports/{record['id']}").status_code == 200
    assert client.delete(f"/api/exports/{record['id']}").status_code == 404
