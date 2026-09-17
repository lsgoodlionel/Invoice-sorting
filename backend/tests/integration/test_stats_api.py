"""统计 API（T14/T15/T16）与统计导出。"""

import io
from urllib.parse import quote

import openpyxl
import pytest

from tests.integration.test_batches_helpers import (
    add_items,
    category_id,
    make_batch,
    make_expense,
    project_id,
)

STATUSES = {"spent", "invoiced", "complete", "sent", "reimbursed", "void"}
Q3 = {"start": "2026-07-01", "end": "2026-09-30"}
Q4 = {"start": "2026-10-01", "end": "2026-12-31"}


def stats(client, **params) -> dict:
    response = client.get("/api/stats", params=params)
    assert response.status_code == 200, response.text
    return response.json()["data"]


def listing(client, **params) -> dict:
    response = client.get("/api/expenses", params={"page_size": 500, **params})
    assert response.status_code == 200, response.text
    return response.json()["data"]


@pytest.fixture
def lifecycle(client):
    """9 月支出 → 10 月外发 → 11 月到账的一条记录，外加一条未外发记录。"""
    travelled = make_expense(client, spent_on="2026-09-20", cents=50000, merchant="滴滴")
    make_expense(client, spent_on="2026-09-25", cents=7000, merchant="文具店", with_invoice=False)
    batch = make_batch(client)
    add_items(client, batch["id"], [travelled["id"]])
    client.post(f"/api/batches/{batch['id']}/sent", json={"sent_on": "2026-10-08"})
    client.post(f"/api/batches/{batch['id']}/received", json={"received_on": "2026-11-10"})
    return travelled


def test_date_basis_assigns_periods(client, lifecycle):  # T14
    by_spent = stats(client, **Q3)
    assert by_spent["date_basis"] == "spent"
    assert by_spent["totals"]["spent_cents"] == 57000
    assert by_spent["totals"]["reimbursed_cents"] == 50000
    assert by_spent["totals"]["pending_cents"] == 7000
    assert by_spent["months"] == [
        {"month": "2026-07", "amount_cents": 0},
        {"month": "2026-08", "amount_cents": 0},
        {"month": "2026-09", "amount_cents": 57000},
    ]
    assert stats(client, **Q3, date_basis="received")["totals"]["spent_cents"] == 0
    by_received = stats(client, **Q4, date_basis="received")
    assert by_received["totals"]["spent_cents"] == 50000
    assert [m["amount_cents"] for m in by_received["months"]] == [0, 50000, 0]
    by_sent = stats(client, **Q4, date_basis="sent", group_by="month")
    assert [(row["key"], row["total_cents"]) for row in by_sent["rows"]] == [("2026-10", 50000)]
    assert stats(client, **Q4, date_basis="invoiced")["rows"] == []  # 图片发票无开票日期


def test_void_is_listed_separately(client):  # T15
    kept = make_expense(client, cents=10000, merchant="保留")
    void = make_expense(client, cents=2500, merchant="退货")
    client.post(f"/api/expenses/{void['id']}/status", json={"status": "void", "note": "退货"})
    data = stats(client, **Q3)
    assert data["totals"]["spent_cents"] == kept["amount_cents"]
    assert data["totals"]["void_cents"] == 2500
    row = data["rows"][0]
    assert (row["key"], row["label"], row["total_cents"]) == ("none", "未分类", 10000)
    assert row["by_status"]["void"] == {"count": 1, "amount_cents": 2500}
    assert set(row["by_status"]) == STATUSES
    assert data["months"][-1]["amount_cents"] == 10000


def test_drilldown_matches_expense_list(client):  # T16
    consumables = category_id(client, "易耗品")
    batch = make_batch(client)
    sent = [
        make_expense(client, cents=96000 + i, merchant=f"京东{i}", category_id=consumables)
        for i in range(3)
    ]
    make_expense(client, cents=1234, merchant="未外发", category_id=consumables)
    make_expense(
        client, spent_on="2026-10-02", cents=999, merchant="区间外", category_id=consumables
    )
    add_items(client, batch["id"], [e["id"] for e in sent], force=True)
    client.post(f"/api/batches/{batch['id']}/sent", json={"sent_on": "2026-09-28"})

    for basis in ("spent", "sent"):
        data = stats(client, **Q3, date_basis=basis)
        row = next(r for r in data["rows"] if r["label"] == "易耗品")
        assert row["key"] == str(consumables)
        for status in STATUSES:
            listed = listing(client, **Q3, date_basis=basis, category_id=consumables, status=status)
            assert listed["total_cents"] == row["by_status"][status]["amount_cents"]
            assert listed["total"] == row["by_status"][status]["count"]
        everything = listing(client, **Q3, date_basis=basis)
        assert (
            everything["total_cents"]
            == data["totals"]["spent_cents"] + data["totals"]["void_cents"]
        )


def test_group_by_project_and_merchant(client):
    project = project_id(client, "科研A")
    make_expense(client, cents=350, merchant="甲", project_id=project)
    make_expense(client, cents=200, merchant="乙")
    make_expense(client, cents=100, merchant="甲")
    by_project = stats(client, **Q3, group_by="project")["rows"]
    assert [(r["key"], r["label"], r["total_cents"]) for r in by_project] == [
        (str(project), "科研A", 350),
        ("none", "无项目", 300),
    ]
    by_merchant = stats(client, **Q3, group_by="merchant")["rows"]
    assert [(r["key"], r["total_cents"]) for r in by_merchant] == [("甲", 450), ("乙", 200)]


def test_stats_validation(client):
    assert client.get("/api/stats", params={"start": "2026-09-01"}).status_code == 422
    reversed_range = client.get("/api/stats", params={"start": "2026-09-02", "end": "2026-09-01"})
    assert reversed_range.status_code == 400
    assert reversed_range.json()["error"] == "开始日期不能晚于结束日期"
    too_long = client.get("/api/stats", params={"start": "1990-01-01", "end": "2026-09-01"})
    assert too_long.status_code == 400
    bad_group = client.get("/api/stats", params={**Q3, "group_by": "status"})
    assert bad_group.status_code == 422


def test_stats_export_xlsx(client, lifecycle):
    response = client.get("/api/stats/export", params={**Q3, "group_by": "month"})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    expected = quote("统计_2026-07-01_2026-09-30.xlsx")
    assert f"filename*=UTF-8''{expected}" in response.headers["content-disposition"]
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["汇总", "明细"]
    summary = list(workbook["汇总"].iter_rows(values_only=True))
    assert "按支出日期统计" in summary[0][0]
    assert summary[1][:2] == ("支出总额（不含作废）", 570)
    assert summary[7][:2] == ("月份", "合计（不含作废）")
    assert summary[8][:2] == ("2026-09", 570)
    details = list(workbook["明细"].iter_rows(values_only=True))
    assert len(details) == 3
    assert details[1][1:4] == ("2026-09-20", "2026-09-20", "滴滴")
    assert details[1][7:] == ("已报销", 500, 500)
