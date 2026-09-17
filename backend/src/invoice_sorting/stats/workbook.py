"""统计导出 XLSX：“汇总” sheet（总额 + 分组 × 状态）与“明细” sheet。"""

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from invoice_sorting.common.constants import EXPENSE_STATUS_LABELS, ExpenseStatus
from invoice_sorting.exporter.workbook import AMOUNT_FORMAT, cents_to_decimal
from invoice_sorting.stats.service import DATE_BASIS_LABELS, StatsEntry, StatsQuery

BOLD = Font(bold=True)
TOTAL_LABELS = (
    ("spent_cents", "支出总额（不含作废）"),
    ("pending_cents", "未外发（已支出/已开票/凭证齐全）"),
    ("in_transit_cents", "在途（已外发）"),
    ("reimbursed_cents", "已报销（到账金额）"),
    ("void_cents", "个人承担/作废"),
)
GROUP_LABELS = {"category": "分类", "project": "经费项目", "merchant": "商家", "month": "月份"}
DETAIL_HEADERS = (
    "ID", "口径日期", "支出日期", "商家", "摘要", "分类", "经费项目", "状态", "金额", "到账金额"
)  # fmt: skip
DETAIL_AMOUNT_COLUMNS = (9, 10)
COLUMN_WIDTH = 16


def _bold_row(sheet: Worksheet) -> None:
    for cell in sheet[sheet.max_row]:
        cell.font = BOLD


def _format_amounts(sheet: Worksheet, columns: range | tuple[int, ...]) -> None:
    for column in columns:
        sheet.cell(row=sheet.max_row, column=column).number_format = AMOUNT_FORMAT


def _write_summary(sheet: Worksheet, query: StatsQuery, stats: dict[str, Any]) -> None:
    sheet.title = "汇总"
    basis = DATE_BASIS_LABELS[query.date_basis]
    sheet.append((f"统计区间 {query.start} ~ {query.end}（按{basis}统计）",))
    sheet["A1"].font = Font(bold=True, size=13)
    for key, label in TOTAL_LABELS:
        sheet.append((label, cents_to_decimal(stats["totals"][key])))
        _format_amounts(sheet, (2,))
    sheet.append(())
    statuses = list(ExpenseStatus)
    header = [GROUP_LABELS[str(query.group_by)], "合计（不含作废）"]
    header += [EXPENSE_STATUS_LABELS[status] for status in statuses]
    sheet.append(header)
    _bold_row(sheet)
    for row in stats["rows"]:
        amounts = [row["by_status"][status.value]["amount_cents"] for status in statuses]
        sheet.append([row["label"], *(cents_to_decimal(c) for c in [row["total_cents"], *amounts])])
        _format_amounts(sheet, range(2, len(header) + 1))
    for index in range(1, len(header) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = COLUMN_WIDTH
    sheet.column_dimensions["A"].width = COLUMN_WIDTH * 2


def _write_details(sheet: Worksheet, entries: list[StatsEntry]) -> None:
    sheet.append(DETAIL_HEADERS)
    _bold_row(sheet)
    sheet.freeze_panes = "A2"
    ordered = sorted(entries, key=lambda entry: (entry.basis_date, entry.expense.id))
    for entry in ordered:
        expense = entry.expense
        sheet.append(
            (
                expense.id,
                entry.basis_date.isoformat(),
                expense.spent_on.isoformat(),
                expense.merchant,
                expense.summary,
                expense.category.name if expense.category else "",
                expense.project.name if expense.project else "",
                EXPENSE_STATUS_LABELS.get(ExpenseStatus(expense.status), expense.status),
                cents_to_decimal(expense.amount_cents),
                cents_to_decimal(expense.reimbursed_cents or 0),
            )
        )
        _format_amounts(sheet, DETAIL_AMOUNT_COLUMNS)
    for index in range(1, len(DETAIL_HEADERS) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = COLUMN_WIDTH


def build_stats_workbook(
    query: StatsQuery, stats: dict[str, Any], entries: list[StatsEntry]
) -> bytes:
    workbook = Workbook()
    _write_summary(workbook.active, query, stats)
    _write_details(workbook.create_sheet("明细"), entries)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
