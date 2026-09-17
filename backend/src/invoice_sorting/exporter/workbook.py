"""资料包汇总表：Sheet“汇总”与“分类小计”。金额以分累加，写入时转 Decimal 元。"""

import io
from collections.abc import Callable
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from invoice_sorting.attachments.serializers import kind_label
from invoice_sorting.common.constants import ChecklistLevel, ChecklistState
from invoice_sorting.db.models import Batch, Expense
from invoice_sorting.expenses.serializers import first_invoice_no
from invoice_sorting.exporter.naming import PackageItem, kind_rank

AMOUNT_FORMAT = "#,##0.00"
LIST_SEPARATOR = "、"
NO_CATEGORY = "未分类"
SUMMARY_HEADERS = (
    "序号",
    "日期",
    "商家",
    "摘要",
    "分类",
    "金额",
    "发票号码",
    "附件清单",
    "缺项",
    "备注",
)
SUMMARY_WIDTHS = (6, 12, 28, 24, 12, 14, 24, 36, 24, 36)
AMOUNT_COLUMN = SUMMARY_HEADERS.index("金额") + 1
SUBTOTAL_HEADERS = ("分类", "条数", "金额")
SUBTOTAL_WIDTHS = (16, 8, 16)
HEADER_ROW = 2
BOLD = Font(bold=True)

NoteLookup = Callable[[int], str]


def cents_to_decimal(cents: int) -> Decimal:
    """分 → 元（Decimal，两位小数），避免浮点误差。"""
    return Decimal(cents).scaleb(-2)


def attachment_list(expense: Expense) -> str:
    kinds = sorted({attachment.kind for attachment in expense.attachments}, key=kind_rank)
    return LIST_SEPARATOR.join(kind_label(kind) for kind in kinds)


def missing_list(expense: Expense) -> str:
    missing = [
        item.attachment_kind
        for item in expense.checklist_items
        if item.level == ChecklistLevel.REQUIRED and item.state == ChecklistState.MISSING
    ]
    return LIST_SEPARATOR.join(kind_label(kind) for kind in sorted(missing, key=kind_rank))


def _note(expense: Expense, extra: str) -> str:
    return "；".join(part for part in (expense.note or "", extra) if part)


def _set_widths(sheet: Worksheet, widths: tuple[int, ...]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _write_header(sheet: Worksheet, title: str, headers: tuple[str, ...]) -> None:
    sheet.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=column, value=header)
        cell.font = BOLD
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = f"A{HEADER_ROW + 1}"


def _summary_row(item: PackageItem, extra_note: str) -> tuple:
    expense = item.expense
    return (
        item.seq,
        expense.spent_on.isoformat(),
        expense.merchant,
        expense.summary,
        expense.category.name if expense.category else NO_CATEGORY,
        cents_to_decimal(expense.amount_cents),
        first_invoice_no(expense) or "",
        attachment_list(expense),
        missing_list(expense),
        _note(expense, extra_note),
    )


def _write_summary(
    sheet: Worksheet, title: str, items: list[PackageItem], notes: NoteLookup
) -> None:
    sheet.title = "汇总"
    _write_header(sheet, title, SUMMARY_HEADERS)
    for item in items:
        sheet.append(_summary_row(item, notes(item.expense.id)))
        sheet.cell(row=sheet.max_row, column=AMOUNT_COLUMN).number_format = AMOUNT_FORMAT
    total = sum(item.expense.amount_cents for item in items)
    sheet.append(("合计", "", "", "", f"{len(items)} 条", cents_to_decimal(total)))
    for column in range(1, len(SUMMARY_HEADERS) + 1):
        sheet.cell(row=sheet.max_row, column=column).font = BOLD
    sheet.cell(row=sheet.max_row, column=AMOUNT_COLUMN).number_format = AMOUNT_FORMAT
    _set_widths(sheet, SUMMARY_WIDTHS)


def category_subtotals(items: list[PackageItem]) -> list[tuple[str, int, int]]:
    """(分类, 条数, 金额分)，按金额降序。"""
    groups: dict[str, tuple[int, int]] = {}
    for item in items:
        expense = item.expense
        name = expense.category.name if expense.category else NO_CATEGORY
        count, cents = groups.get(name, (0, 0))
        groups[name] = (count + 1, cents + expense.amount_cents)
    rows = [(name, count, cents) for name, (count, cents) in groups.items()]
    return sorted(rows, key=lambda row: (-row[2], row[0]))


def _write_subtotals(sheet: Worksheet, title: str, items: list[PackageItem]) -> None:
    _write_header(sheet, title, SUBTOTAL_HEADERS)
    for name, count, cents in category_subtotals(items):
        sheet.append((name, count, cents_to_decimal(cents)))
        sheet.cell(row=sheet.max_row, column=3).number_format = AMOUNT_FORMAT
    total = sum(item.expense.amount_cents for item in items)
    sheet.append(("合计", len(items), cents_to_decimal(total)))
    for column in range(1, 4):
        sheet.cell(row=sheet.max_row, column=column).font = BOLD
    sheet.cell(row=sheet.max_row, column=3).number_format = AMOUNT_FORMAT
    _set_widths(sheet, SUBTOTAL_WIDTHS)


def build_summary_workbook(
    batch: Batch, items: list[PackageItem], notes: NoteLookup, moment: datetime
) -> bytes:
    project = batch.project
    project_name = project.name if project else "无项目"
    title = f"批次：{batch.name}　项目：{project_name}　导出时间：{moment:%Y-%m-%d %H:%M}"
    workbook = Workbook()
    _write_summary(workbook.active, title, items, notes)
    _write_subtotals(workbook.create_sheet("分类小计"), title, items)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
